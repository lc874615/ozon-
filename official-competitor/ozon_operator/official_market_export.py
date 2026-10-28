"""Import a user-downloaded Ozon XLSX/CSV; no network or account writes.
Header aliases are exact: unfamiliar columns remain unknown rather than guessed.
"""
import csv
import hashlib
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from .market_selection import evaluate_market_selection


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '').replace('\u00a0', ' ')).strip().casefold().rstrip('.')


ALIASES = {
    'sku': {'sku', 'ozon sku', 'артикул ozon', 'sku ozon', 'ozon商品编号'},
    'brand': {'бренд', 'brand', '品牌'},
    'scheme': {'схема продаж', 'схема работы', 'sales schema', '发货方式', '履约方式'},
    'sold': {'продано, шт', 'продано товаров, шт', 'количество проданных товаров', 'sales, units', '销量'},
    'sold28': {'продано за 28 дней, шт', 'продано товаров за 28 дней, шт', '销量（近28天）', 'sales in 28 days'},
    'ordered': {'заказано товаров, шт', 'заказано, шт', 'ordered units', '订购件数'},
    'ordered28': {'заказано товаров за 28 дней, шт', 'заказано за 28 дней, шт'},
}
ALIASES = {k: {clean(v) for v in values} for k, values in ALIASES.items()}
PACKAGING = {
    'packaging_length_cm': ('length', 'cm'), 'packaging_width_cm': ('width', 'cm'), 'packaging_height_cm': ('height', 'cm'),
    'gross_weight_kg': ('weight', 'kg'),
    'длина упаковки, см': ('length', 'cm'), 'ширина упаковки, см': ('width', 'cm'), 'высота упаковки, см': ('height', 'cm'),
    'длина упаковки, мм': ('length', 'mm'), 'ширина упаковки, мм': ('width', 'mm'), 'высота упаковки, мм': ('height', 'mm'),
    'вес товара в упаковке, г': ('weight', 'g'), 'вес с упаковкой, г': ('weight', 'g'),
    'вес товара в упаковке, кг': ('weight', 'kg'), 'вес с упаковкой, кг': ('weight', 'kg'),
    '包装长度（cm）': ('length', 'cm'), '包装宽度（cm）': ('width', 'cm'), '包装高度（cm）': ('height', 'cm'),
    '毛重（kg）': ('weight', 'kg'),
}


def numeric(value):
    if isinstance(value, bool):
        return None
    try:
        n = Decimal(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
        return n if n.is_finite() and n >= 0 else None
    except InvalidOperation:
        return None


def sku_string(value):
    # IDs must not be reconstructed from scientific notation or rounded floats.
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, str) and re.fullmatch(r'\d+', value.strip()):
        return value.strip()
    return None


def table_rows(path):
    p = Path(path)
    if p.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('File exceeds 10 MiB')
    if p.suffix.lower() == '.csv':
        text = p.read_text(encoding='utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8000], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(text.splitlines(), dialect))
        if len(rows) > 30000:
            raise ValueError('Row limit')
        return [('CSV', rows)]
    if p.suffix.lower() != '.xlsx':
        raise ValueError('Only UTF-8 CSV and XLSX supported')
    with zipfile.ZipFile(p) as z:
        if len(z.infolist()) > 10000 or sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
            raise ValueError('XLSX decompressed size limit')
    from openpyxl import load_workbook
    workbook = load_workbook(p, read_only=True, data_only=True, keep_links=False)
    sheets, count = [], 0
    try:
        for sheet in workbook:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                count += 1
                if count > 30000 or len(row) > 200:
                    raise ValueError('Worksheet size limit')
                rows.append(list(row))
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def import_official_market_export(path, expected_sku, confirmed_period_days=None):
    sku = sku_string(expected_sku)
    if not sku:
        raise ValueError('Expected exact numeric SKU required for single-product test')
    if confirmed_period_days not in (None, 28):
        raise ValueError('User approved official 28-day period only')
    digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    matches, header_maps = [], []
    for sheet, rows in table_rows(path):
        header_index, mapped = None, {}
        for index, row in enumerate(rows[:50]):
            candidate = {}
            for col, value in enumerate(row):
                label = clean(value)
                for key, aliases in ALIASES.items():
                    if label in aliases:
                        if key in candidate:
                            raise ValueError('Duplicate semantic column: ' + key)
                        candidate[key] = col
            if 'sku' in candidate:
                header_index, mapped = index, candidate
                break
        if header_index is None:
            continue
        headers = rows[header_index]
        header_maps.append({'sheet': sheet, 'header_row': header_index + 1,
                            'recognized_columns': {k: str(headers[v]) for k, v in mapped.items()}})
        for offset, row in enumerate(rows[header_index + 1:], header_index + 2):
            def cell(key):
                col = mapped.get(key)
                return row[col] if col is not None and col < len(row) else None
            if sku_string(cell('sku')) != sku:
                continue
            metric, count, chosen_key = None, None, None
            for key, kind in [('sold28', 'sold_units'), ('ordered28', 'ordered_units'), ('sold', 'sold_units'), ('ordered', 'ordered_units')]:
                if key in mapped:
                    n = numeric(cell(key))
                    count = int(n) if n is not None and n == n.to_integral_value() else None
                    metric, chosen_key = kind, key
                    break
            period28 = chosen_key in ('sold28', 'ordered28') or confirmed_period_days == 28
            package_values, package_sources = {}, []
            for col, header in enumerate(headers):
                definition = PACKAGING.get(clean(header))
                if definition is None or col >= len(row):
                    continue
                n = numeric(row[col])
                if n is None or n <= 0:
                    continue
                axis, unit = definition
                package_values.setdefault(axis, set()).add(n / (1000 if unit == 'g' else 10 if unit == 'mm' else 1))
                package_sources.append({'column': str(header), 'cell_column': col + 1, 'unit': unit})
            conflicts = [k for k, values in package_values.items() if len(values) > 1]
            package = {k: str(next(iter(values))) for k, values in package_values.items() if len(values) == 1}
            complete = not conflicts and all(k in package for k in ('length', 'width', 'height', 'weight'))
            signal = {'status': 'observed', 'source_type': 'user_supplied_official_export', 'source_sha256': digest,
                      'sales_period_days': 28 if period28 else None, 'sales_period_verified': period28,
                      'period_evidence': 'explicit_28_day_column' if chosen_key in ('sold28', 'ordered28') else 'user_confirmed_UI_period' if period28 else None,
                      'metric_kind': metric,
                      'items': [{'sku': sku, 'brand': cell('brand'), 'salesSchema': cell('scheme'),
                                 'soldCount': count if metric == 'sold_units' else None}]}
            snapshot = {'product_id': sku, 'official_market_export': signal}
            selection = evaluate_market_selection(snapshot)
            todo = []
            if metric == 'ordered_units':
                todo.append('ordered_units_must_not_be_silently_relabelled_as_sold_units')
            if not complete:
                todo.append('competitor_packaging_missing_no_creation_fallback')
            matches.append({'product_id': sku, 'sheet': sheet, 'row': offset,
                            'official_market_export': signal, 'selection': selection,
                            'observed_metric': {'kind': metric, 'count': count},
                            'packaging': {'complete': complete, 'conflicts': conflicts, 'dimensions_cm': {k: package[k] for k in ('length', 'width', 'height') if k in package},
                                          'gross_kg': package.get('weight'), 'field_sources': package_sources,
                                          'physical_measurements_verified': False},
                            'todo': todo, 'publication_approved': False})
    if len(matches) > 1:
        for item in matches:
            item['selection']['status'] = 'review'
            item['selection']['pending_reasons'].append('multiple_export_rows_for_exact_sku_require_review')
    return {'source_type': 'user_supplied_official_export', 'source_sha256': digest,
            'expected_sku': sku, 'status': 'single_sku_found' if len(matches) == 1 else 'duplicate_sku_requires_review' if matches else 'sku_or_columns_not_found',
            'header_maps': header_maps, 'matched_rows': matches, 'live_network_requests': 0, 'store_mutations': 0,
            'publication_approved': False}
